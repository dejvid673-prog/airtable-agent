from __future__ import annotations

import hashlib
import shutil
from copy import copy
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Finding
from .product_transform import TransformResult


GENERATED_SHEETS = (
    "EXPORT_GOTOWY",
    "RAPORT_AGENTA",
    "DO_WERYFIKACJI",
    "MAPA_KATEGORII",
    "PLAN_ZMIAN",
)


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_sheet(path: str | Path, sheet_name: str) -> tuple[list[str], list[list[Any]], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Brak arkusza: {sheet_name}")
        sheet = workbook[sheet_name]
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            raise ValueError(f"Arkusz {sheet_name} jest pusty.")
        headers = [str(value) if value is not None else "" for value in values[0]]
        rows = [list(row) for row in values[1:] if any(value is not None for value in row)]
        return headers, rows, list(workbook.sheetnames)
    finally:
        workbook.close()


def _replace_sheet(workbook, name: str):
    if name in workbook.sheetnames:
        workbook.remove(workbook[name])
    return workbook.create_sheet(name)


def _write_matrix(sheet, rows: Iterable[Iterable[Any]]) -> None:
    for row in rows:
        sheet.append(list(row))


def _style_header(sheet, columns: int, fill: str = "1F4E78") -> None:
    for cell in sheet[1][:columns]:
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _fit_columns(sheet, max_width: int = 48) -> None:
    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        width = 10
        for cell in column_cells[:250]:
            if cell.value is not None:
                width = max(width, min(len(str(cell.value)) + 2, max_width))
        sheet.column_dimensions[letter].width = width


def _write_findings(sheet, findings: list[Finding], source_headers: list[str], source_rows: list[list[Any]]) -> None:
    header = [
        "Ważność", "Kod", "Wiersz źródłowy", "ID produktu", "SKU", "Nazwa produktu",
        "Pole", "Bieżąca wartość", "Opis", "Zalecenie", "Powiązane wiersze",
    ]
    sheet.append(header)
    source = {index: dict(zip(source_headers, row)) for index, row in enumerate(source_rows, start=2)}
    order = {"BLOCKER": 0, "ERROR": 1, "WARNING": 2, "REVIEW": 3, "INFO": 4}
    for item in sorted(findings, key=lambda value: (order.get(value.severity, 9), value.code, value.row or 0)):
        row = source.get(item.row or -1, {})
        sheet.append([
            item.severity,
            item.code,
            item.row,
            row.get("ID produktu"),
            row.get("SKU"),
            row.get("Nazwa produktu"),
            item.field,
            item.current_value,
            item.message,
            item.suggested_action,
            ", ".join(str(value) for value in item.related_rows),
        ])
    _style_header(sheet, len(header), fill="9C0006")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _fit_columns(sheet)


def write_result_workbook(
    input_path: str | Path,
    output_path: str | Path,
    source_headers: list[str],
    source_rows: list[list[Any]],
    result: TransformResult,
    contract: dict[str, Any],
) -> dict[str, Any]:
    source = Path(input_path)
    output = Path(output_path)
    if source.resolve() == output.resolve():
        raise ValueError("Plik wynikowy musi mieć inną ścieżkę niż wejściowy.")
    if source.suffix.lower() != ".xlsx" or output.suffix.lower() != ".xlsx":
        raise ValueError("Wersja 0.2.0 obsługuje pliki .xlsx.")

    output.parent.mkdir(parents=True, exist_ok=True)
    before_hash = sha256_file(source)
    shutil.copy2(source, output)
    workbook = load_workbook(output)
    try:
        prepared = _replace_sheet(workbook, contract["prepared_sheet"])
        _write_matrix(prepared, [result.headers, *result.rows])
        _style_header(prepared, len(result.headers))
        _fit_columns(prepared)
        number_formats = {
            "Waga produktu kg": "0.000",
            "Waga wysyłkowa kg": "0.000",
            "Wytrzymałość kg": "0.000",
            "Cena netto": "0.00",
            "Stan": "0",
            "EAN": "@",
        }
        header_index = {cell.value: cell.column for cell in prepared[1]}
        for name, number_format in number_formats.items():
            column = header_index.get(name)
            if column:
                for row in range(2, prepared.max_row + 1):
                    prepared.cell(row=row, column=column).number_format = number_format

        report = _replace_sheet(workbook, "RAPORT_AGENTA")
        report_rows = [
            ["Metryka", "Wartość"],
            ["Wersja kontraktu", contract["contract_version"]],
            ["Arkusz źródłowy", contract["source_sheet"]],
            ["Liczba rekordów", result.metrics.get("records", 0)],
            ["Wszystkie ustalenia", result.metrics.get("findings", 0)],
            ["Nazwy ujednolicone", result.metrics.get("names_normalized", 0)],
            ["Wagi produktu wydzielone", result.metrics.get("product_weight_extracted", 0)],
            ["Długości wydzielone", result.metrics.get("length_extracted", 0)],
            ["Pojemności wydzielone", result.metrics.get("volume_extracted", 0)],
            ["Ilości wydzielone", result.metrics.get("quantity_extracted", 0)],
            ["Wagi wysyłkowe dodane", result.metrics.get("shipping_weights_added", 0)],
            ["SHA-256 wejścia", before_hash],
            ["Zasada", "Arkusze źródłowe zachowano; wynik zapisano w nowych arkuszach."],
        ]
        for name, count in sorted(result.metrics.get("main_category_counts", {}).items()):
            report_rows.append([f"Kategoria: {name}", count])
        _write_matrix(report, report_rows)
        _style_header(report, 2)
        _fit_columns(report)

        review = _replace_sheet(workbook, "DO_WERYFIKACJI")
        _write_findings(review, result.findings, source_headers, source_rows)

        category_map = _replace_sheet(workbook, "MAPA_KATEGORII")
        _write_matrix(category_map, [["Kategoria główna", "Podkategoria", "Liczba produktów"], *result.category_map])
        _style_header(category_map, 3, fill="548235")
        _fit_columns(category_map)

        plan = _replace_sheet(workbook, "PLAN_ZMIAN")
        _write_matrix(plan, [
            ["Operacja", "Status", "Zasada"],
            ["Ujednolicenie nazw", "WYKONANO", "Tylko bezpieczne porządkowanie spacji i wydzielenie segmentów po |."],
            ["Wydzielenie długości/pojemności/ilości/wagi", "WYKONANO", "Wartości trafiają do osobnych kolumn."],
            ["Mapowanie kategorii głównych", "WYKONANO", "Tylko według jawnego category_mapping."],
            ["Usuwanie lub scalanie rekordów", "ZABLOKOWANO", "Wymaga osobnej decyzji człowieka."],
            ["Zmiana cen, stanów, aktywności i EAN", "ZABLOKOWANO", "Brak zgody w kontrakcie."],
        ])
        _style_header(plan, 3, fill="548235")
        _fit_columns(plan)

        workbook.save(output)
    finally:
        workbook.close()

    after_hash = sha256_file(source)
    if after_hash != before_hash:
        raise RuntimeError("Plik wejściowy został zmodyfikowany.")
    verification = verify_result_workbook(source, output, contract)
    return {
        "input_sha256": before_hash,
        "output_sha256": sha256_file(output),
        "verification": verification,
    }


def verify_result_workbook(input_path: str | Path, output_path: str | Path, contract: dict[str, Any]) -> dict[str, Any]:
    input_headers, input_rows, _ = read_sheet(input_path, contract["source_sheet"])
    output_headers, output_rows, output_sheets = read_sheet(output_path, contract["prepared_sheet"])
    input_index = {name: pos for pos, name in enumerate(input_headers)}
    output_index = {name: pos for pos, name in enumerate(output_headers)}
    checks = {
        "input_exists": Path(input_path).exists(),
        "output_exists": Path(output_path).exists(),
        "row_count_equal": len(input_rows) == len(output_rows),
        "product_ids_equal": [row[input_index["ID produktu"]] for row in input_rows] == [row[output_index["ID produktu"]] for row in output_rows],
        "skus_equal": [row[input_index["SKU"]] for row in input_rows] == [row[output_index["SKU"]] for row in output_rows],
        "generated_sheets_present": all(name in output_sheets for name in GENERATED_SHEETS),
        "main_categories_valid": all(row[output_index["Kategoria główna"]] in contract["allowed_main_categories"] for row in output_rows),
    }
    return {
        "passed": all(checks.values()),
        "checks": checks,
        "input_rows": len(input_rows),
        "output_rows": len(output_rows),
    }
