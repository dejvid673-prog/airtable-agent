import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from airtable_workbook_agent.contracts import load_workbook_contract
from airtable_workbook_agent.product_transform import transform_rows
from airtable_workbook_agent.xlsx_backend import read_sheet, sha256_file, verify_result_workbook, write_result_workbook


class XlsxBackendTests(unittest.TestCase):
    def test_full_local_run_preserves_source(self):
        contract = load_workbook_contract(Path(__file__).parents[1] / "contracts" / "product_workbook_contract.json")
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "input.xlsx"
            output = root / "output.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "EXPORT"
            sheet.append(contract["required_columns"])
            sheet.append([1, "SKU-1", "Karma dla ryb | 1 kg", "Ryby i rośliny", "Karma dla ryb", "1 kg", None, 10.0, 5, True, None])
            helper = workbook.create_sheet("DANE_POMOCNICZE")
            helper["A1"] = "test"
            workbook.save(source)
            before = sha256_file(source)
            headers, rows, _ = read_sheet(source, "EXPORT")
            result = transform_rows(headers, rows, contract)
            execution = write_result_workbook(source, output, headers, rows, result, contract)
            self.assertEqual(before, sha256_file(source))
            self.assertTrue(execution["verification"]["passed"])
            self.assertTrue(output.exists())
            saved = load_workbook(output, read_only=True)
            try:
                for name in ["EXPORT_GOTOWY", "RAPORT_AGENTA", "DO_WERYFIKACJI", "MAPA_KATEGORII", "PLAN_ZMIAN"]:
                    self.assertIn(name, saved.sheetnames)
                self.assertIn("DANE_POMOCNICZE", saved.sheetnames)
            finally:
                saved.close()
            verify = verify_result_workbook(source, output, contract)
            self.assertTrue(verify["passed"])


if __name__ == "__main__":
    unittest.main()
